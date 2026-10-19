#!/usr/bin/env python3
"""
Load the AKI/AKB pilot tables from a Profil Kesehatan workbook into PostgreSQL.

Source mapping (sheet number -> fact table)
    21  births                        -> fact.fact_birth
    23  maternal deaths by phase      -> fact.fact_maternal_death
    25  maternal deaths by cause      -> fact.fact_maternal_death_cause
    34  child deaths by age and sex   -> fact.fact_child_death
    36  neonatal deaths by cause      -> fact.fact_child_death_cause
    37  child (12-59 month) causes    -> fact.fact_child_death_cause

Regions come from table 21 (kecamatan -> puskesmas), causes from the header
rows of tables 25/36/37, and the year/district from the title block.

Every load is idempotent: rows are upserted on their natural key and tagged
with the import_batch_id. A QA section then reconciles counts loaded into the
database against the workbook's own TOTAL rows and the Resume sheet.

Usage:
    python scripts/load_profilkes.py profilkes_2025.xlsx
    DATABASE_URL=postgresql://user:pass@host:5432/db python scripts/load_profilkes.py file.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import os
from collections import OrderedDict
from pathlib import Path

import openpyxl
import psycopg

DEFAULT_DSN = "postgresql://postgres:postgres@localhost:55432/profilkes"

SHEETS = {
    "birth": "21",
    "maternal_death": "23",
    "maternal_cause": "25",
    "child_death": "34",
    "neonatal_cause": "36",
    "child_cause": "37",
}

# The workbook leaves the kecamatan cell blank for these three puskesmas in
# every one of its 88 tables (verified), and no cell anywhere states it. Each
# sits inside the block of the kecamatan named here, which also matches where
# the puskesmas is located, so assign it explicitly instead of guessing.
KECAMATAN_OVERRIDES = {
    "WIDOROPAYUNG": "BESUKI",
    "KLAMPOKAN": "PANJI",
    "WONOREJO": "BANYUPUTIH",
}

# Not present in the workbook: every title only names the district. Override
# with --province/--province-bps if you load a different province.
PROVINCE_NAME = "JAWA TIMUR"
PROVINCE_BPS = "35"
DISTRICT_BPS = "3512"  # Kabupaten Situbondo

TOTAL_PREFIX = "TOTAL"
TOTAL_ALT_PREFIX = "JUMLAH (KAB/KOTA)"


class DataError(RuntimeError):
    """Raised when the workbook does not look like the expected template."""


# ---------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------
def clean(value) -> str:
    """Collapse newlines/extra spaces so header and label matching is stable."""
    return " ".join(str(value).split())


def upper(value) -> str:
    return clean(value).upper()


def num(value, ctx: str, warnings: list[str]) -> int:
    """Coerce a count cell to int. Empty cells are 0 and are recorded."""
    if value is None or (isinstance(value, str) and not value.strip()):
        warnings.append(f"{ctx}: empty cell read as 0")
        return 0
    if isinstance(value, bool):
        raise DataError(f"{ctx}: boolean where a count was expected")
    if isinstance(value, (int, float)):
        if float(value) != int(value):
            raise DataError(f"{ctx}: non-integer count {value!r}")
        return int(value)
    raise DataError(f"{ctx}: non-numeric count {value!r}")


def is_total_row(row) -> bool:
    """True for the trailing summary rows, e.g. 'TOTAL' or 'JUMLAH (KAB/KOTA)'."""
    first = row[0] if row else None
    if not isinstance(first, str):
        return False
    key = upper(first)
    return key.startswith(TOTAL_PREFIX) or key.startswith("JUMLAH (")


def cell(row, idx: int):
    return row[idx] if len(row) > idx else None


def iter_data_rows(rows, pkm_index, sheet: str):
    """Yield (sheet_row_number, puskesmas_UPPER, row) for numeric data rows."""
    for row_no, row in enumerate(rows, 1):
        if is_total_row(row):
            return
        first = row[0] if row else None
        if not isinstance(first, (int, float)):
            continue
        # The template has a numbering row (1, 2, 3, ...) between the header
        # and the data; it has no text in the first three columns.
        if not any(isinstance(cell(row, i), str) for i in range(3)):
            continue
        pkm = None
        # Column C normally holds the puskesmas name; some tables put it in B
        # and leave a literal 0 in C, so fall back to the kecamatan column.
        for idx in (2, 1):
            value = cell(row, idx)
            if isinstance(value, str) and upper(value) in pkm_index:
                pkm = upper(value)
                break
        if pkm is None:
            raise DataError(
                f"sheet {sheet} row {row_no}: cannot resolve puskesmas from {row[:3]!r}"
            )
        yield row_no, pkm, row


def find_total_row(rows, labels: tuple[str, ...]):
    for row in rows:
        first = row[0] if row else None
        if not isinstance(first, str):
            continue
        key = upper(first)
        if any(key.startswith(label) for label in labels):
            return row
    return None


def scan_causes(rows, first_cause_prefix: str):
    """
    Find the cause header row and return ([(col_idx, name), ...], total_col_idx).

    Cause labels live on one row, start at the first known cause and stop at the
    first blank cell or the 'JUMLAH ...' total column.
    """
    for row in rows:
        for idx, value in enumerate(row):
            if isinstance(value, str) and upper(value).startswith(first_cause_prefix):
                causes: list[tuple[int, str]] = []
                for c in range(idx, len(row)):
                    label = cell(row, c)
                    if not isinstance(label, str) or not clean(label):
                        break
                    if upper(label).startswith("JUMLAH"):
                        return causes, c
                    causes.append((c, clean(label)))
                return causes, None
    raise DataError(f"cause header starting with {first_cause_prefix!r} not found")


def find_label_value(rows, labels: tuple[str, ...], max_rows: int = 10):
    """Return the first non-empty cell to the right of a label cell."""
    for row in rows[:max_rows]:
        for idx, value in enumerate(row):
            if isinstance(value, str) and upper(value) in labels:
                for nxt in row[idx + 1:]:
                    if nxt not in (None, ""):
                        return nxt
    return None


def resume_number(resume, label: str):
    """Read a Resume indicator: L+P column first, else the first numeric of L/P/L+P."""
    for row in resume:
        if isinstance(cell(row, 1), str) and upper(cell(row, 1)) == label:
            total = cell(row, 4)
            if isinstance(total, (int, float)) and not isinstance(total, bool):
                return total
            for value in row[2:5]:
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    return value
            return None
    return None


# ---------------------------------------------------------------------
# hierarchy
# ---------------------------------------------------------------------
def read_hierarchy(birth_rows, warnings: list[str]):
    """Return OrderedDict puskesmas -> kecamatan, in workbook order."""
    pkm_to_kec: "OrderedDict[str, str]" = OrderedDict()
    last_kec: str | None = None
    for row_no, row in enumerate(birth_rows, 1):
        if is_total_row(row):
            break
        if not isinstance(cell(row, 0), (int, float)):
            continue
        if not any(isinstance(cell(row, i), str) for i in range(3)):
            continue
        pkm_cell = cell(row, 2)
        if not isinstance(pkm_cell, str) or not clean(pkm_cell):
            continue
        pkm = upper(pkm_cell)
        kec_cell = cell(row, 1)
        kec = upper(kec_cell) if isinstance(kec_cell, str) and clean(kec_cell) else None

        if kec:
            source = "stated"
        elif pkm in KECAMATAN_OVERRIDES:
            kec, source = KECAMATAN_OVERRIDES[pkm], "override"
        elif last_kec:
            kec, source = last_kec, "inferred"
            warnings.append(
                f"sheet {SHEETS['birth']} row {row_no}: kecamatan for {pkm} inferred as {kec}"
            )
        else:
            raise DataError(f"sheet {SHEETS['birth']} row {row_no}: no kecamatan for {pkm}")

        if source != "stated":
            print(f"  kecamatan {kec:<14} for {pkm:<14} ({source})")
        last_kec = kec
        pkm_to_kec[pkm] = kec

    if not pkm_to_kec:
        raise DataError(f"no puskesmas rows found in table {SHEETS['birth']}")
    return pkm_to_kec


def upsert_region(cur, level: str, parent_id, name: str, bps_code=None) -> int:
    cur.execute(
        """
        INSERT INTO dim.dim_region (parent_id, level, name, bps_code)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (level, parent_id, name) DO UPDATE
           SET bps_code = COALESCE(EXCLUDED.bps_code, dim.dim_region.bps_code)
        RETURNING region_id
        """,
        (parent_id, level, name, bps_code),
    )
    return cur.fetchone()[0]


def upsert_causes(cur, cause_group: str, causes) -> dict[str, int]:
    ids: dict[str, int] = {}
    for sort_order, (_, name) in enumerate(causes, start=1):
        cur.execute(
            """
            INSERT INTO dim.dim_cause (cause_group, name, sort_order)
            VALUES (%s, %s, %s)
            ON CONFLICT (cause_group, name) DO UPDATE SET sort_order = EXCLUDED.sort_order
            RETURNING cause_id
            """,
            (cause_group, name, sort_order),
        )
        ids[name] = cur.fetchone()[0]
    return ids


# ---------------------------------------------------------------------
# load
# ---------------------------------------------------------------------
def upsert_facts(cur, table: str, columns: list[str], conflict: str, rows: list[tuple]):
    if not rows:
        return
    placeholders = ", ".join(["%s"] * len(columns))
    updates = ", ".join(f"{c} = EXCLUDED.{c}" for c in columns if c not in conflict.split(", "))
    sql = (
        f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders}) "
        f"ON CONFLICT ({conflict}) DO UPDATE SET {updates}"
    )
    cur.executemany(sql, rows)
    print(f"  {table:<38} {len(rows):>3} rows")


def load(path: Path, dsn: str, province: str, province_bps: str, district_bps: str) -> None:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    warnings: list[str] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    rows = {key: [list(r) for r in wb[name].iter_rows(values_only=True)]
            for key, name in SHEETS.items()}
    resume = [list(r) for r in wb["Resume"].iter_rows(values_only=True)]

    district = find_label_value(rows["maternal_death"], ("KABUPATEN/KOTA", "KABUPATEN"))
    year = find_label_value(rows["maternal_death"], ("TAHUN",))
    if not district:
        raise DataError("district name not found in the title block of table 23")
    if not isinstance(year, (int, float)):
        raise DataError(f"reporting year not found, got {year!r}")
    district, year = upper(district), int(year)
    if district != "SITUBONDO" and district_bps == DISTRICT_BPS:
        print(f"  district {district} != SITUBONDO: ignoring default BPS code")
        district_bps = ""

    print(f"Source   : {path.name} (sha256 {digest[:12]}...)")
    print(f"Reporting: {district} {year}\n")
    print("Regions")
    pkm_to_kec = read_hierarchy(rows["birth"], warnings)

    with psycopg.connect(dsn) as conn, conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO etl.import_batch (source_file, file_sha256, source_year, district)
            VALUES (%s, %s, %s, %s) RETURNING batch_id
            """,
            (path.name, digest, year, district),
        )
        batch_id = cur.fetchone()[0]

        province_id = upsert_region(cur, "province", None, province, province_bps or None)
        district_id = upsert_region(cur, "district", province_id, district, district_bps or None)

        kec_ids: dict[str, int] = {}
        region_ids: dict[str, int] = {}
        for pkm, kec in pkm_to_kec.items():
            if kec not in kec_ids:
                kec_ids[kec] = upsert_region(cur, "subdistrict", district_id, kec)
            region_ids[pkm] = upsert_region(cur, "health_center", kec_ids[kec], pkm)

        cur.execute(
            "INSERT INTO dim.dim_period (year) VALUES (%s) "
            "ON CONFLICT (year) DO UPDATE SET year = EXCLUDED.year RETURNING period_id",
            (year,),
        )
        period_id = cur.fetchone()[0]
        print(f"  {len(kec_ids)} kecamatan, {len(region_ids)} puskesmas, period {year}\n")

        # ---- births (21) ------------------------------------------------
        print("Facts")
        birth_facts = []
        for row_no, pkm, row in iter_data_rows(rows["birth"], region_ids, SHEETS["birth"]):
            ctx = f"t21 r{row_no}"
            m_live, m_still = num(cell(row, 3), ctx, warnings), num(cell(row, 4), ctx, warnings)
            f_live, f_still = num(cell(row, 6), ctx, warnings), num(cell(row, 7), ctx, warnings)
            lp_live, lp_still = num(cell(row, 9), ctx, warnings), num(cell(row, 10), ctx, warnings)
            if m_live + f_live != lp_live or m_still + f_still != lp_still:
                raise DataError(
                    f"{ctx}: sex split {m_live}+{f_live}/{m_still}+{f_still} "
                    f"does not match total {lp_live}/{lp_still}"
                )
            birth_facts += [
                (region_ids[pkm], period_id, "M", m_live, m_still, batch_id),
                (region_ids[pkm], period_id, "F", f_live, f_still, batch_id),
            ]
        upsert_facts(
            cur, "fact.fact_birth",
            ["region_id", "period_id", "sex", "live_births", "stillbirths", "import_batch_id"],
            "region_id, period_id, sex", birth_facts,
        )

        # ---- maternal deaths by phase (23) ------------------------------
        phase_facts = []
        for row_no, pkm, row in iter_data_rows(rows["maternal_death"], region_ids, SHEETS["maternal_death"]):
            ctx = f"t23 r{row_no}"
            values = [num(cell(row, i), ctx, warnings) for i in (3, 4, 5)]
            reported = num(cell(row, 6), ctx, warnings)
            if sum(values) != reported:
                raise DataError(f"{ctx}: phases {values} do not sum to reported total {reported}")
            for phase, deaths in zip(("pregnant", "delivery", "postpartum"), values):
                phase_facts.append((region_ids[pkm], period_id, phase, deaths, batch_id))
        upsert_facts(
            cur, "fact.fact_maternal_death",
            ["region_id", "period_id", "phase", "deaths", "import_batch_id"],
            "region_id, period_id, phase", phase_facts,
        )

        # ---- maternal deaths by cause (25) ------------------------------
        causes_25, total_25 = scan_causes(rows["maternal_cause"], "KOMPLIKASI ABORTUS")
        cause_ids_maternal = upsert_causes(cur, "maternal", causes_25)
        maternal_cause_facts = []
        for row_no, pkm, row in iter_data_rows(rows["maternal_cause"], region_ids, SHEETS["maternal_cause"]):
            ctx = f"t25 r{row_no}"
            deaths = [num(cell(row, c), ctx, warnings) for c, _ in causes_25]
            if total_25 is not None:
                reported = num(cell(row, total_25), ctx, warnings)
                if sum(deaths) != reported:
                    raise DataError(f"{ctx}: causes sum {sum(deaths)} != reported {reported}")
            for (_, name), value in zip(causes_25, deaths):
                maternal_cause_facts.append(
                    (region_ids[pkm], period_id, cause_ids_maternal[name], value, batch_id)
                )
        upsert_facts(
            cur, "fact.fact_maternal_death_cause",
            ["region_id", "period_id", "cause_id", "deaths", "import_batch_id"],
            "region_id, period_id, cause_id", maternal_cause_facts,
        )

        # ---- child deaths by age and sex (34) ---------------------------
        # M block cols 4-8, F cols 9-13, combined cols 14-18
        # (neonatal, post neonatal, bayi, anak balita, total) in each block
        child_facts = []
        for row_no, pkm, row in iter_data_rows(rows["child_death"], region_ids, SHEETS["child_death"]):
            ctx = f"t34 r{row_no}"
            m = [num(cell(row, i), ctx, warnings) for i in range(3, 8)]
            f = [num(cell(row, i), ctx, warnings) for i in range(8, 13)]
            c = [num(cell(row, i), ctx, warnings) for i in range(13, 18)]
            for label, block in (("M", m), ("F", f)):
                if block[2] != block[0] + block[1]:
                    raise DataError(f"{ctx}: {label} bayi {block[2]} != neonatal+post {block[0] + block[1]}")
                if block[4] != block[2] + block[3]:
                    raise DataError(f"{ctx}: {label} total {block[4]} != bayi+balita {block[2] + block[3]}")
            for i in range(5):
                if m[i] + f[i] != c[i]:
                    raise DataError(f"{ctx}: combined {c[i]} != M+F {m[i] + f[i]} at position {i}")
            child_facts += [
                (region_ids[pkm], period_id, "M", "neonatal", m[0], batch_id),
                (region_ids[pkm], period_id, "M", "post_neonatal", m[1], batch_id),
                (region_ids[pkm], period_id, "M", "child_1_4", m[3], batch_id),
                (region_ids[pkm], period_id, "F", "neonatal", f[0], batch_id),
                (region_ids[pkm], period_id, "F", "post_neonatal", f[1], batch_id),
                (region_ids[pkm], period_id, "F", "child_1_4", f[3], batch_id),
            ]
        upsert_facts(
            cur, "fact.fact_child_death",
            ["region_id", "period_id", "sex", "age_band", "deaths", "import_batch_id"],
            "region_id, period_id, sex, age_band", child_facts,
        )

        # ---- child deaths by cause (36 neonatal, 37 child) --------------
        cause_facts = []
        for key, group, age_band, first in (
            ("neonatal_cause", "neonatal", "neonatal", "MALFORMASI KONGENITAL"),
            ("child_cause", "child", "child_1_4", "DIARE"),
        ):
            causes, _ = scan_causes(rows[key], first)
            cause_ids = upsert_causes(cur, group, causes)
            for row_no, pkm, row in iter_data_rows(rows[key], region_ids, SHEETS[key]):
                ctx = f"{SHEETS[key]} r{row_no}"
                for col, name in causes:
                    cause_facts.append(
                        (region_ids[pkm], period_id, cause_ids[name], age_band,
                         num(cell(row, col), ctx, warnings), batch_id)
                    )
        upsert_facts(
            cur, "fact.fact_child_death_cause",
            ["region_id", "period_id", "cause_id", "age_band", "deaths", "import_batch_id"],
            "region_id, period_id, cause_id", cause_facts,
        )

        # ---- QA ---------------------------------------------------------
        print("\nReconciliation")
        compare(rows, resume, cur, year, district, warnings)

    wb.close()

    if warnings:
        print(f"\n{len(warnings)} warning(s):")
        for message in warnings[:12]:
            print(f"  - {message}")
        if len(warnings) > 12:
            print(f"  ... {len(warnings) - 12} more")


def compare(rows, resume, cur, year: int, district: str, warnings: list[str]) -> None:
    """Reconcile DB contents against the workbook's own totals and Resume."""
    totals = {}
    total21 = find_total_row(rows["birth"], ("TOTAL", "JUMLAH ("))
    totals["live_births"] = num(cell(total21, 9), "t21 total", warnings)
    totals["stillbirths"] = num(cell(total21, 10), "t21 total", warnings)

    total23 = find_total_row(rows["maternal_death"], ("TOTAL",))
    totals["maternal_deaths"] = num(cell(total23, 6), "t23 total", warnings)
    totals["deaths_pregnant"] = num(cell(total23, 3), "t23 total", warnings)
    totals["deaths_delivery"] = num(cell(total23, 4), "t23 total", warnings)
    totals["deaths_postpartum"] = num(cell(total23, 5), "t23 total", warnings)

    causes_25, _ = scan_causes(rows["maternal_cause"], "KOMPLIKASI ABORTUS")
    total25 = find_total_row(rows["maternal_cause"], ("TOTAL",))
    totals["maternal_cause_deaths"] = sum(
        num(cell(total25, col), "t25 total", warnings) for col, _ in causes_25
    )

    total34 = find_total_row(rows["child_death"], ("TOTAL",))
    totals["neonatal_deaths"] = num(cell(total34, 13), "t34 total", warnings)
    totals["post_neonatal_deaths"] = num(cell(total34, 14), "t34 total", warnings)
    totals["infant_deaths"] = num(cell(total34, 15), "t34 total", warnings)
    totals["child_1_4_deaths"] = num(cell(total34, 16), "t34 total", warnings)
    totals["under5_deaths"] = num(cell(total34, 17), "t34 total", warnings)

    causes_36, _ = scan_causes(rows["neonatal_cause"], "MALFORMASI KONGENITAL")
    total36 = find_total_row(rows["neonatal_cause"], ("TOTAL",))
    totals["neonatal_cause_deaths"] = sum(
        num(cell(total36, col), "t36 total", warnings) for col, _ in causes_36
    )
    causes_37, _ = scan_causes(rows["child_cause"], "DIARE")
    total37 = find_total_row(rows["child_cause"], ("TOTAL",))
    totals["child_cause_deaths"] = sum(
        num(cell(total37, col), "t37 total", warnings) for col, _ in causes_37
    )

    actual = dict(
        zip(
            ("live_births", "stillbirths", "maternal_deaths", "deaths_pregnant",
             "deaths_delivery", "deaths_postpartum", "maternal_cause_deaths",
             "neonatal_deaths", "post_neonatal_deaths", "infant_deaths",
             "child_1_4_deaths", "under5_deaths", "neonatal_cause_deaths",
             "child_cause_deaths"),
            cur.execute(
                """
                SELECT (SELECT SUM(live_births) FROM fact.fact_birth),
                       (SELECT SUM(stillbirths) FROM fact.fact_birth),
                       (SELECT SUM(deaths) FROM fact.fact_maternal_death),
                       (SELECT SUM(deaths) FROM fact.fact_maternal_death WHERE phase = 'pregnant'),
                       (SELECT SUM(deaths) FROM fact.fact_maternal_death WHERE phase = 'delivery'),
                       (SELECT SUM(deaths) FROM fact.fact_maternal_death WHERE phase = 'postpartum'),
                       (SELECT SUM(deaths) FROM fact.fact_maternal_death_cause),
                       (SELECT SUM(deaths) FROM fact.fact_child_death WHERE age_band = 'neonatal'),
                       (SELECT SUM(deaths) FROM fact.fact_child_death WHERE age_band = 'post_neonatal'),
                       (SELECT SUM(deaths) FROM fact.fact_child_death
                         WHERE age_band IN ('neonatal', 'post_neonatal')),
                       (SELECT SUM(deaths) FROM fact.fact_child_death WHERE age_band = 'child_1_4'),
                       (SELECT SUM(deaths) FROM fact.fact_child_death),
                       (SELECT SUM(d.deaths) FROM fact.fact_child_death_cause d
                          JOIN dim.dim_cause c ON c.cause_id = d.cause_id
                         WHERE c.cause_group = 'neonatal'),
                       (SELECT SUM(d.deaths) FROM fact.fact_child_death_cause d
                          JOIN dim.dim_cause c ON c.cause_id = d.cause_id
                         WHERE c.cause_group = 'child')
                """
            ).fetchone(),
        )
    )

    kpi = cur.execute(
        """
        SELECT maternal_deaths, live_births, aki_per_100k,
               neonatal_deaths, infant_deaths, under5_deaths, akb_per_1000,
               aki_ci95_low, aki_ci95_high, akb_ci95_low, akb_ci95_high
        FROM mart.kpi_summary
        WHERE region_name = %s AND year = %s
        """,
        (district, year),
    ).fetchone()
    (k_maternal, k_live, k_aki, k_neonatal, k_infant, k_under5, k_akb,
     aki_lo, aki_hi, akb_lo, akb_hi) = kpi

    failures = 0

    def check(name: str, expected, got, tolerance: float = 0) -> None:
        nonlocal failures
        ok = expected is not None and got is not None and abs(float(expected) - float(got)) <= tolerance
        failures += 0 if ok else 1
        shown_expected = "-" if expected is None else f"{expected:g}"
        shown_got = "-" if got is None else f"{got:g}"
        print(f"  {'PASS' if ok else 'FAIL'}  {name:<34} source={shown_expected:<10} db={shown_got}")

    for key, label in (
        ("live_births", "live births"),
        ("stillbirths", "stillbirths"),
        ("maternal_deaths", "maternal deaths"),
        ("deaths_pregnant", "  during pregnancy"),
        ("deaths_delivery", "  during delivery"),
        ("deaths_postpartum", "  postpartum"),
        ("maternal_cause_deaths", "maternal deaths by cause"),
        ("neonatal_deaths", "neonatal deaths"),
        ("post_neonatal_deaths", "post-neonatal deaths"),
        ("infant_deaths", "infant deaths"),
        ("child_1_4_deaths", "child 12-59 month deaths"),
        ("under5_deaths", "under-5 deaths"),
        ("neonatal_cause_deaths", "neonatal deaths by cause"),
        ("child_cause_deaths", "child deaths by cause"),
    ):
        check(label, totals.get(key), actual.get(key))

    # Rates are derived, so check them against the Resume where it is usable.
    check("Resume live births", resume_number(resume, "JUMLAH LAHIR HIDUP"), k_live)
    check("Resume maternal deaths", resume_number(resume, "JUMLAH KEMATIAN IBU"), k_maternal)
    check("Resume infant deaths", resume_number(resume, "JUMLAH BAYI MATI"), k_infant)
    check("Resume AKI (published)", resume_number(resume, "ANGKA KEMATIAN IBU (DILAPORKAN)"), k_aki, 0.01)
    print(
        f"  INFO  AKI derived {k_aki}/100k (95% CI {aki_lo}-{aki_hi}), "
        f"AKB derived {k_akb}/1000 (95% CI {akb_lo}-{akb_hi})"
    )
    print(
        "  INFO  Resume AKB rate is #DIV/0! in the source, so only the 102-infant "
        "count can be reconciled there; the rate above is computed from counts."
    )
    if k_neonatal != actual.get("neonatal_deaths") or k_under5 != actual.get("under5_deaths"):
        failures += 1
        print("  FAIL  KPI view disagrees with fact totals")

    if failures:
        raise SystemExit(f"\n{failures} reconciliation check(s) failed.")
    print("\nAll reconciliation checks passed.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--dsn", default=os.environ.get("DATABASE_URL", DEFAULT_DSN))
    parser.add_argument("--province", default=PROVINCE_NAME)
    parser.add_argument("--province-bps", default=PROVINCE_BPS)
    parser.add_argument("--district-bps", default=DISTRICT_BPS)
    args = parser.parse_args()

    if not args.workbook.exists():
        raise SystemExit(f"workbook not found: {args.workbook}")
    load(args.workbook, args.dsn, upper(args.province), args.province_bps, args.district_bps)


if __name__ == "__main__":
    main()
